"""P1 — static-file extraction adapters (Word/Excel/PDF → markdown)."""
import io

import pytest

from integrations.extract import ExtractionError, extract


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Product Requirements", level=1)
    doc.add_paragraph("The system shall let a user log in.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "Title"
    table.rows[1].cells[0].text = "US-1"
    table.rows[1].cells[1].text = "Login"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Stories"
    ws.append(["ID", "Title"])
    ws.append(["US-1", "Login"])
    ws.append(["US-2", "Logout"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_docx_headings_paragraphs_tables():
    title, body = extract(filename="PRD.docx", mime="", content=_docx_bytes())
    assert title == "PRD"
    assert "# Product Requirements" in body
    assert "The system shall let a user log in." in body
    assert "| ID | Title |" in body and "| US-1 | Login |" in body


def test_extract_xlsx_sheets_as_tables():
    title, body = extract(filename="backlog.xlsx", mime="", content=_xlsx_bytes())
    assert title == "backlog"
    assert "## Stories" in body
    assert "| ID | Title |" in body
    assert "US-1" in body and "US-2" in body


def test_extract_txt_returns_body_as_is():
    title, body = extract(
        filename="context.txt", mime="text/plain",
        content=b"# Notes\n\nThe billing service handles invoices.\n",
    )
    assert title == "context"
    assert "# Notes" in body
    assert "The billing service handles invoices." in body


def test_extract_html_strips_tags_keeps_structure():
    html = (
        b"<html><head><style>.x{color:red}</style><title>t</title></head>"
        b"<body><h1>Design</h1><p>Screens for <b>login</b>.</p>"
        b"<ul><li>Home</li><li>Profile</li></ul>"
        b"<script>alert('x')</script></body></html>"
    )
    title, body = extract(filename="ux.html", mime="text/html", content=html)
    assert title == "ux"
    assert "# Design" in body
    assert "Screens for login." in body
    assert "- Home" in body and "- Profile" in body
    # script/style contents are dropped
    assert "alert" not in body and "color:red" not in body


def test_unsupported_type_raises():
    with pytest.raises(ExtractionError):
        extract(filename="logo.png", mime="image/png", content=b"\x89PNG\r\n")


def test_corrupt_pdf_wrapped_as_extraction_error():
    # routes to the pdf adapter (by extension), which fails on garbage bytes.
    with pytest.raises(ExtractionError):
        extract(filename="broken.pdf", mime="application/pdf", content=b"not a real pdf")
