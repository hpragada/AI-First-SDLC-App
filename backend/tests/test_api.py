"""P0-T11 — REST API surface.

Exercises the HTTP layer end-to-end against the real domain services, stubbing only
the Temporal-touching seam (`_default_workflow_starter` / `_default_signal_sender`) at
module level so Run/Approval rows are genuinely created. Tenant scoping is asserted
directly (the test connection is superuser, so RLS is bypassed — the explicit
`tenant_id` filter is what's under test here; RLS itself is covered by test_rls.py).
"""
import pytest
from rest_framework.test import APIClient

from apps.gates import services as gate_services
from apps.gates.models import Approval
from apps.projects.models import Project, ProjectProfile
from apps.registry.catalog import seed_agent_specs
from apps.registry.models import AgentSpec, Artifact, TraceEdge
from apps.runs.models import Run
from apps.tenancy.models import Membership, Tenant, User
from apps.workitems.models import WorkItem
from orchestration import trigger as trigger_mod


@pytest.fixture
def client():
    return APIClient()


def _hdr(tenant_id):
    return {"HTTP_X_TENANT_ID": str(tenant_id)}


@pytest.fixture
def tenant(db):
    return Tenant.objects.create(name="Acme")


@pytest.fixture
def project(tenant):
    return Project.objects.create(tenant=tenant, name="Platform")


# ── Tenant bootstrap ────────────────────────────────────────────────
def test_create_tenant(client, db):
    resp = client.post("/api/tenants/", {"name": "NewCo"}, format="json")
    assert resp.status_code == 201
    assert resp.data["name"] == "NewCo"
    tenant = Tenant.objects.get(id=resp.data["id"], name="NewCo")
    assert Tenant.objects.filter(id=tenant.id, name="NewCo").exists()
    assert AgentSpec.objects.filter(tenant=tenant).count() > 0
    assert AgentSpec.objects.filter(tenant=tenant, key="01_ideation").exists()


# ── Projects + profile ──────────────────────────────────────────────
def test_create_and_list_projects(client, tenant):
    resp = client.post("/api/projects/", {"name": "Apollo"}, format="json", **_hdr(tenant.id))
    assert resp.status_code == 201
    pid = resp.data["id"]
    assert Project.objects.get(id=pid).tenant_id == tenant.id

    listed = client.get("/api/projects/", **_hdr(tenant.id))
    assert listed.status_code == 200
    assert [p["id"] for p in listed.data] == [pid]


def test_profile_upsert_bumps_version(client, project):
    url = f"/api/projects/{project.id}/profile/"
    # GET before any profile → 404
    assert client.get(url, **_hdr(project.tenant_id)).status_code == 404

    created = client.put(
        url, {"important_context": "north-star"}, format="json", **_hdr(project.tenant_id)
    )
    assert created.status_code == 201
    assert created.data["important_context"] == "north-star"
    assert created.data["version"] == 1

    updated = client.put(
        url, {"important_context": "revised"}, format="json", **_hdr(project.tenant_id)
    )
    assert updated.status_code == 200
    assert updated.data["version"] == 2
    assert ProjectProfile.objects.get(project=project).important_context == "revised"


# ── WorkItems ───────────────────────────────────────────────────────
def test_create_and_filter_work_items(client, project):
    resp = client.post(
        "/api/work-items/",
        {"project": str(project.id), "type": "story", "feature_stream": "auth"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 201
    wid = resp.data["id"]
    assert WorkItem.objects.get(id=wid).state == WorkItem.State.PROPOSED

    by_project = client.get(f"/api/work-items/?project={project.id}", **_hdr(project.tenant_id))
    assert [w["id"] for w in by_project.data] == [wid]
    # State filter that excludes it returns empty.
    none = client.get("/api/work-items/?state=done", **_hdr(project.tenant_id))
    assert none.data == []


def test_create_work_item_rejects_foreign_project(client, tenant):
    other = Tenant.objects.create(name="Other")
    foreign = Project.objects.create(tenant=other, name="Theirs")
    resp = client.post(
        "/api/work-items/", {"project": str(foreign.id)}, format="json", **_hdr(tenant.id)
    )
    assert resp.status_code == 400
    assert "project" in resp.data


# ── Phase trigger (real trigger_phase, stubbed workflow start) ──────
def test_trigger_phase_creates_run(client, project, monkeypatch):
    seed_agent_specs(project.tenant)
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    started = []
    monkeypatch.setattr(
        trigger_mod, "_default_workflow_starter",
        lambda wf_id, inp: started.append((wf_id, inp.phase)),
    )

    resp = client.post(
        f"/api/work-items/{wi.id}/trigger/", {"agent_key": "01_ideation"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 202
    run = Run.objects.get(id=resp.data["run_id"])
    assert run.work_item_id == wi.id
    assert run.temporal_workflow_id == resp.data["workflow_id"]
    assert started and started[0][1] == "ideation"


def test_trigger_unknown_agent_is_400(client, project, monkeypatch):
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    monkeypatch.setattr(trigger_mod, "_default_workflow_starter", lambda *a, **k: None)
    resp = client.post(
        f"/api/work-items/{wi.id}/trigger/", {"agent_key": "nope"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 400


# ── Runs / artifacts / events / trace-edges (read) ──────────────────
def test_list_runs_artifacts_events_edges(client, project):
    spec = AgentSpec.objects.create(tenant=project.tenant, key="03_prd", body_ref="x")
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    run = Run.objects.create(tenant=project.tenant, work_item=wi, agent_spec=spec)
    a1 = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="requirements", type="requirements"
    )
    a2 = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="prd", type="prd"
    )
    TraceEdge.objects.create(tenant=project.tenant, from_artifact=a1, to_artifact=a2, rel="derives")

    runs = client.get(f"/api/runs/?project={project.id}", **_hdr(project.tenant_id))
    assert [r["id"] for r in runs.data] == [str(run.id)]

    arts = client.get(f"/api/artifacts/?project={project.id}&phase=prd", **_hdr(project.tenant_id))
    assert [a["id"] for a in arts.data] == [str(a2.id)]

    edges = client.get(f"/api/trace-edges/?project={project.id}", **_hdr(project.tenant_id))
    assert edges.data[0]["rel"] == "derives"

    events = client.get(f"/api/events/?target_id={run.id}", **_hdr(project.tenant_id))
    assert events.status_code == 200


def test_traceability_endpoint_returns_nodes_and_edges(client, project):
    run = Run.objects.create(
        tenant=project.tenant, work_item=WorkItem.objects.create(
            tenant=project.tenant, project=project, type=WorkItem.Type.STORY),
        agent_spec=AgentSpec.objects.create(tenant=project.tenant, key="03_prd", body_ref="x"),
    )
    req = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="requirements",
        type="requirements", status="published", sor="internal",
    )
    prd = Artifact.objects.create(
        tenant=project.tenant, project=project, run=run, phase="prd", type="prd",
        status="published", sor="confluence",
        external_ref={"sor": "confluence", "id": "123", "url": "https://x/wiki/123", "version": 1},
    )
    TraceEdge.objects.create(
        tenant=project.tenant, from_artifact=req, to_artifact=prd, rel="derives"
    )

    resp = client.get(f"/api/projects/{project.id}/traceability/", **_hdr(project.tenant_id))
    assert resp.status_code == 200
    nodes = {n["id"]: n for n in resp.data["nodes"]}
    assert set(nodes) == {str(req.id), str(prd.id)}
    # The PRD node carries its live SoR deep-link.
    assert nodes[str(prd.id)]["sor"] == "confluence"
    assert nodes[str(prd.id)]["url"] == "https://x/wiki/123"
    assert nodes[str(req.id)]["url"] is None
    assert resp.data["edges"] == [
        {"from": str(req.id), "to": str(prd.id), "rel": "derives"}
    ]


# ── Gate decision (real decide_gate, stubbed signal) ────────────────
@pytest.fixture
def gate(project, monkeypatch):
    spec = AgentSpec.objects.create(tenant=project.tenant, key="01_ideation", body_ref="x")
    wi = WorkItem.objects.create(tenant=project.tenant, project=project, type=WorkItem.Type.STORY)
    run = Run.objects.create(
        tenant=project.tenant, work_item=wi, agent_spec=spec, temporal_workflow_id="wf-x"
    )
    t = project.tenant
    po = User.objects.create_user(email="po@acme.test", name="PO", tenant=t, password="p")
    Membership.objects.create(tenant=t, user=po, project=project, role="product_owner")
    eng = User.objects.create_user(email="e@acme.test", name="E", tenant=t, password="p")
    Membership.objects.create(tenant=t, user=eng, project=project, role="engineer")
    approval = Approval.objects.create(
        tenant=project.tenant, run=run, work_item=wi, required_role="product_owner",
    )
    signals = []
    monkeypatch.setattr(
        gate_services, "_default_signal_sender",
        lambda wf, dec, uid, comment: signals.append((wf, dec)),
    )
    return {"approval": approval, "po": po, "eng": eng, "signals": signals}


def test_decide_gate_approved(client, project, gate):
    resp = client.post(
        f"/api/approvals/{gate['approval'].id}/decide/",
        {"user_id": str(gate["po"].id), "decision": "approved", "comment": "lgtm"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 200
    gate["approval"].refresh_from_db()
    assert gate["approval"].decision == Approval.Decision.APPROVED
    assert gate["signals"] == [("wf-x", "approved")]


def test_decide_gate_wrong_role_is_403(client, project, gate):
    resp = client.post(
        f"/api/approvals/{gate['approval'].id}/decide/",
        {"user_id": str(gate["eng"].id), "decision": "approved"},
        format="json", **_hdr(project.tenant_id),
    )
    assert resp.status_code == 403
    gate["approval"].refresh_from_db()
    assert gate["approval"].decision == Approval.Decision.PENDING
    assert gate["signals"] == []


def test_list_pending_approvals(client, project, gate):
    resp = client.get("/api/approvals/?decision=pending", **_hdr(project.tenant_id))
    assert [a["id"] for a in resp.data] == [str(gate["approval"].id)]


# ── Brownfield-lite onboarding ──────────────────────────────────────
def _docx_bytes():
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("PRD", level=1)
    doc.add_paragraph("The system shall let a user log in.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_source_document_upload_extract_and_adopt(client, project, monkeypatch):
    from django.core.files.uploadedfile import SimpleUploadedFile

    from integrations import onboarding
    from storage import InMemoryObjectStore

    monkeypatch.setattr(onboarding, "_object_store", lambda injected: InMemoryObjectStore())
    po = User.objects.create_user(email="po@x.test", name="PO", tenant=project.tenant, password="p")
    Membership.objects.create(tenant=project.tenant, user=po, project=project, role="product_owner")

    upload = SimpleUploadedFile(
        "PRD.docx", _docx_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    up = client.post(
        f"/api/projects/{project.id}/source-documents/",
        {"file": upload, "artifact_type": "prd"}, format="multipart", **_hdr(project.tenant_id),
    )
    assert up.status_code == 201
    sd_id, aid = up.data["source_document_id"], up.data["artifact_ids"][0]
    assert Artifact.objects.get(id=aid).status == "draft"  # awaiting the review gate

    listed = client.get(
        f"/api/projects/{project.id}/source-documents/", **_hdr(project.tenant_id)
    )
    assert [s["id"] for s in listed.data] == [sd_id]

    # Review gate: product_owner adopts → draft becomes published.
    res = client.post(
        f"/api/source-documents/{sd_id}/adopt/", {"user_id": str(po.id)},
        format="json", **_hdr(project.tenant_id),
    )
    assert res.status_code == 200
    assert Artifact.objects.get(id=aid).status == "published"


def test_onboard_atlassian_endpoint(client, project, monkeypatch):
    from api import views

    monkeypatch.setattr(
        views, "register_atlassian",
        lambda **kw: {"registered_pages": ["a1"], "registered_issues": ["ST-1"]},
    )
    res = client.post(
        f"/api/projects/{project.id}/onboard-atlassian/", {}, format="json",
        **_hdr(project.tenant_id),
    )
    assert res.status_code == 200
    assert res.data == {"registered_pages": ["a1"], "registered_issues": ["ST-1"]}


# ── Tenant scoping ──────────────────────────────────────────────────
def test_tenant_isolation_hides_other_tenants_rows(client, project):
    other = Tenant.objects.create(name="Other")
    listed = client.get("/api/projects/", **_hdr(other.id))
    assert listed.data == []
    # Direct retrieve of a foreign project is a 404 under the other tenant.
    retr = client.get(f"/api/projects/{project.id}/", **_hdr(other.id))
    assert retr.status_code == 404


def test_missing_tenant_header_is_400(client):
    assert client.get("/api/projects/").status_code == 400
