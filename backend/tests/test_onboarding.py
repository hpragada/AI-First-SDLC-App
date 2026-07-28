"""P1 — brownfield-lite onboarding: static-file extraction + review gate, and
Atlassian register-don't-recreate."""
import io

import pytest
from mcp_platform_state import tools

from apps.integrations.models import SourceDocument
from apps.projects.models import Project
from apps.registry.models import Artifact, TraceEdge
from apps.tenancy.models import Membership, Tenant, User
from apps.workitems.models import WorkItem
from integrations.extract import ExtractionError
from integrations.onboarding import (
    OnboardingRoleError,
    adopt_source_document,
    build_gap_report,
    ensure_onboarding_work_item,
    ingest_source_document,
    link_confluence_pages,
    read_gap_report,
    register_atlassian,
    submit_agent_gap_report,
    write_gap_report_artifact,
)
from integrations.types import DocPage, Issue, IssueHit
from orchestration.trigger import NoCompanionError, trigger_onboarding_agent
from storage import InMemoryObjectStore


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("PRD", level=1)
    doc.add_paragraph("The system shall let a user log in.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def world(db):
    t = Tenant.objects.create(name="Acme")
    project = Project.objects.create(tenant=t, name="Legacy App")
    po = User.objects.create_user(email="po@acme.test", name="PO", tenant=t, password="p")
    Membership.objects.create(tenant=t, user=po, project=project, role="product_owner")
    eng = User.objects.create_user(email="e@acme.test", name="E", tenant=t, password="p")
    Membership.objects.create(tenant=t, user=eng, project=project, role="engineer")
    return {"t": t, "project": project, "po": po, "eng": eng, "store": InMemoryObjectStore()}


# ── static-file extraction → draft artifact → review gate ───────────
def test_ingest_creates_draft_artifact_with_provenance(world):
    out = ingest_source_document(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id),
        filename="PRD.docx", mime="", content=_docx_bytes(),
        artifact_type="prd", object_store=world["store"],
    )
    assert out["status"] == "extracted"

    src = SourceDocument.objects.get(id=out["source_document_id"])
    assert src.status == SourceDocument.Status.EXTRACTED
    assert src.provenance["filename"] == "PRD.docx"

    artifact = Artifact.objects.get(id=out["artifact_ids"][0])
    assert artifact.status == Artifact.Status.DRAFT       # awaiting review (#5)
    assert artifact.sor == Artifact.SoR.INTERNAL and artifact.type == "prd"
    assert artifact.external_ref["source"] == "static_file"
    body = world["store"].get(artifact.s3_key).decode("utf-8")
    assert "# PRD" in body and "log in" in body           # extracted, not the original bytes


def test_adopt_publishes_after_role_check_and_is_idempotent(world):
    out = ingest_source_document(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id),
        filename="PRD.docx", mime="", content=_docx_bytes(), object_store=world["store"],
    )
    res = adopt_source_document(
        tenant_id=str(world["t"].id), source_document_id=out["source_document_id"],
        user_id=str(world["po"].id),
    )
    assert res["status"] == "adopted"
    assert Artifact.objects.get(id=out["artifact_ids"][0]).status == Artifact.Status.PUBLISHED
    assert SourceDocument.objects.get(id=out["source_document_id"]).status == \
        SourceDocument.Status.ADOPTED

    again = adopt_source_document(
        tenant_id=str(world["t"].id), source_document_id=out["source_document_id"],
        user_id=str(world["po"].id),
    )
    assert again.get("idempotent") is True


def test_adopt_wrong_role_is_blocked(world):
    out = ingest_source_document(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id),
        filename="PRD.docx", mime="", content=_docx_bytes(), object_store=world["store"],
    )
    with pytest.raises(OnboardingRoleError):
        adopt_source_document(
            tenant_id=str(world["t"].id), source_document_id=out["source_document_id"],
            user_id=str(world["eng"].id),
        )
    assert Artifact.objects.get(id=out["artifact_ids"][0]).status == Artifact.Status.DRAFT


def test_ingest_unsupported_marks_failed(world):
    with pytest.raises(ExtractionError):
        ingest_source_document(
            tenant_id=str(world["t"].id), project_id=str(world["project"].id),
            filename="logo.png", mime="image/png", content=b"\x89PNG", object_store=world["store"],
        )
    assert SourceDocument.objects.get(project=world["project"]).status == \
        SourceDocument.Status.FAILED


def test_ingest_carries_title_onto_artifact(world):
    """Brownfield B1-2: the extracted document title is threaded onto the artifact (was blank)."""
    out = ingest_source_document(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id),
        filename="PRD.docx", mime="", content=_docx_bytes(), object_store=world["store"],
    )
    assert out["title"]
    assert Artifact.objects.get(id=out["artifact_ids"][0]).title == out["title"]


def test_adopt_with_target_type_reclassifies_so_gap_report_sees_it(world):
    """Brownfield B2-1 / B5-1 (headline): an uploaded doc is typed `document`; adopting it AS
    `requirements` reclassifies + publishes it so the gap report counts it as requirements
    evidence. Before the fix the doc was invisible and every phase read 'missing'."""
    tid, pid = str(world["t"].id), str(world["project"].id)
    out = ingest_source_document(
        tenant_id=tid, project_id=pid, filename="Reqs.docx", mime="",
        content=_docx_bytes(), object_store=world["store"],  # default artifact_type -> "document"
    )
    assert Artifact.objects.get(id=out["artifact_ids"][0]).type == "document"

    # Blind before classification: requirements has no matching-typed evidence.
    before = build_gap_report(tenant_id=tid, project_id=pid, object_store=world["store"])
    assert next(p for p in before["phases"] if p["phase"] == "requirements")["status"] == "missing"

    res = adopt_source_document(
        tenant_id=tid, source_document_id=out["source_document_id"],
        user_id=str(world["po"].id), target_type="requirements",
    )
    assert res["target_type"] == "requirements"
    art = Artifact.objects.get(id=out["artifact_ids"][0])
    assert art.type == "requirements" and art.status == Artifact.Status.PUBLISHED

    # Now the gap report recognises it as complete evidence for requirements.
    after = build_gap_report(tenant_id=tid, project_id=pid, object_store=world["store"])
    req = next(p for p in after["phases"] if p["phase"] == "requirements")
    assert req["status"] == "complete"
    assert out["artifact_ids"][0] in [e["artifact_id"] for e in req["evidence"]]


def test_adopt_rejects_invalid_target_type(world):
    """Brownfield B2-1: only canonical document-authored types may be chosen at adopt."""
    from integrations.onboarding import OnboardingError

    out = ingest_source_document(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id),
        filename="PRD.docx", mime="", content=_docx_bytes(), object_store=world["store"],
    )
    with pytest.raises(OnboardingError):
        adopt_source_document(
            tenant_id=str(world["t"].id), source_document_id=out["source_document_id"],
            user_id=str(world["po"].id), target_type="not_a_phase",
        )


# ── Atlassian register-don't-recreate ───────────────────────────────
class _FakeDocStore:
    """B3-4: Confluence is no longer crawled — pages are resolved by id via get_page for typed
    linking. `pages` maps page id → DocPage."""

    def __init__(self, pages):
        self._pages = pages

    def get_page(self, *, page_ref):
        return self._pages[page_ref]


class _FakeTaskTracker:
    def __init__(self, hits, issues):
        self._hits = hits
        self._issues = issues

    def search(self, *, jql):
        return self._hits

    def get_issue(self, *, issue_ref):
        return self._issues[issue_ref]


def test_register_atlassian_imports_jira_backlog_idempotently(world):
    tasktracker = _FakeTaskTracker(
        [IssueHit(key="ST-1", summary="Login"), IssueHit(key="ST-2", summary="Logout")],
        {"ST-1": Issue(key="ST-1", summary="Login", status="To Do", issue_type="Story",
                       url="https://x/browse/ST-1"),
         "ST-2": Issue(key="ST-2", summary="Logout", status="To Do", issue_type="Task",
                       url="https://x/browse/ST-2")},
    )
    kw = dict(tenant_id=str(world["t"].id), project_id=str(world["project"].id),
              tasktracker=tasktracker, project_key="ST")

    out = register_atlassian(**kw)
    assert out["registered_issues"] == ["ST-1", "ST-2"] and out["tasktracker_connected"] is True

    # Issues → accepted WorkItems (register-don't-recreate; B3-4 no longer crawls Confluence).
    issues = WorkItem.objects.filter(project=world["project"], external_ref__isnull=False)
    assert issues.count() == 2 and {w.external_ref["id"] for w in issues} == {"ST-1", "ST-2"}
    assert all(w.state == WorkItem.State.ACCEPTED for w in issues)

    # Register-don't-recreate: a second run adds nothing.
    again = register_atlassian(**kw)
    assert again["registered_issues"] == []
    assert WorkItem.objects.filter(project=world["project"]).count() == 2


def test_register_atlassian_skips_gracefully_when_jira_not_connected(world):
    # B3-1: no tasktracker connected → graceful skip (no error), so onboarding never dead-ends.
    out = register_atlassian(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id))
    assert out == {"registered_issues": [], "tasktracker_connected": False}
    assert WorkItem.objects.filter(project=world["project"]).count() == 0


def _docpage(pid, title):
    return DocPage(id=pid, title=title, body_html="<p/>", version=1,
                   url=f"https://x/wiki/spaces/BILL/pages/{pid}")


def test_link_confluence_pages_registers_typed_pointers_idempotently(world):
    # B3-4: user-provided links, each typed → typed Confluence-pointer artifacts (no body copied).
    docstore = _FakeDocStore({"100": _docpage("100", "Billing PRD"),
                              "200": _docpage("200", "Arch Overview")})
    out = link_confluence_pages(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id), docstore=docstore,
        links=[{"page": "https://x/wiki/spaces/BILL/pages/100/Billing-PRD", "type": "prd"},
               {"page": "200", "type": "architecture"}],
    )
    assert len(out["linked"]) == 2 and out["errors"] == []

    arts = list(Artifact.objects.filter(project=world["project"], sor="confluence"))
    by_type = {a.type: a for a in arts}
    assert set(by_type) == {"prd", "architecture"}
    assert all(a.status == "published" for a in arts)
    assert by_type["prd"].external_ref["id"] == "100"
    assert by_type["prd"].external_ref["title"] == "Billing PRD"

    # Idempotent by page id: re-linking updates the declared type, never duplicates.
    again = link_confluence_pages(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id), docstore=docstore,
        links=[{"page": "100", "type": "requirements"}],
    )
    assert len(again["linked"]) == 1 and again["linked"][0]["created"] is False
    arts2 = list(Artifact.objects.filter(project=world["project"], sor="confluence"))
    assert len(arts2) == 2
    a100 = next(a for a in arts2 if (a.external_ref or {}).get("id") == "100")
    assert a100.type == "requirements"


def test_link_confluence_bad_input_returns_per_row_errors(world):
    docstore = _FakeDocStore({})
    out = link_confluence_pages(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id), docstore=docstore,
        links=[{"page": "https://x/wiki/no-page-id-here", "type": "prd"},   # unparseable link
               {"page": "100", "type": "code"}],                            # non-adoptable type
    )
    assert out["linked"] == [] and len(out["errors"]) == 2
    assert Artifact.objects.filter(project=world["project"], sor="confluence").count() == 0


def test_code_intel_summary_renders_a_deep_readable_report():
    # B4 enrichment: the persisted code-intel artifact shows the analyzer's real depth (the five
    # layers), rendered readably — not just raw counts / a Python-dict repr.
    from integrations.onboarding import _render_code_intel_summary

    body = _render_code_intel_summary("acme/api", {
        "file_count": 1284, "commits": 412, "authors": 7,
        "languages": {".py": 900, ".ts": 384},
        "communities": 5, "top_modules": ["core/app.py"], "entry_points": ["main.py"],
        "hotspot_files": [
            {"file": "billing/charge.py", "score": 88, "churn": 40, "complexity": 22}],
        "dead_code": [{"file": "legacy/unused.py"}],
        "decisions": [{"title": "Move to Stripe", "ref": "ADR-004"}],
    }, "acme/api overview")
    assert "## Risk hotspots" in body and "billing/charge.py" in body
    assert "## Architecture" in body and "core/app.py" in body
    assert "## Likely dead code" in body and "legacy/unused.py" in body
    assert "## Decisions found" in body and "ADR-004" in body
    assert "Python (900)" in body and "TypeScript (384)" in body   # readable, not a raw dict
    assert "{'.py'" not in body                                    # never a bare dict repr


def test_code_intel_summary_degrades_for_a_thin_payload():
    from integrations.onboarding import _render_code_intel_summary

    body = _render_code_intel_summary(
        "acme/api", {"file_count": 10, "commits": 3, "hotspots": 2, "languages": {".py": 10}}, "")
    assert "Files analysed:** 10" in body and "Hotspots:** 2" in body
    assert "Python (10)" in body


def test_linked_confluence_page_feeds_gap_report(world):
    # B3-4 headline (Q-1 sibling for Confluence): a typed linked page is real gap-report evidence,
    # so the resume point moves off the phase it covers — the whole point of typed linking.
    docstore = _FakeDocStore({"300": _docpage("300", "Requirements")})
    link_confluence_pages(
        tenant_id=str(world["t"].id), project_id=str(world["project"].id), docstore=docstore,
        links=[{"page": "300", "type": "requirements"}],
    )
    report = build_gap_report(tenant_id=str(world["t"].id), project_id=str(world["project"].id),
                              object_store=world["store"])
    status = {p["phase"]: p["status"] for p in report["phases"]}
    assert status["requirements"] == "complete"
    assert report["resume_phase"] != "requirements"


# ── full onboarding: reconcile → gap report → resume (D7) ───────────
def _published(world, type_):
    art = tools.register_artifact(
        str(world["t"].id), str(world["project"].id), None,
        phase="onboarding", type=type_, sor="internal", content_hash="h")
    tools.update_artifact(str(world["t"].id), art["id"], status="published")
    return art["id"]


def _draft(world, type_):
    return tools.register_artifact(
        str(world["t"].id), str(world["project"].id), None,
        phase="onboarding", type=type_, sor="internal", content_hash="h")["id"]


def test_gap_report_empty_project_resumes_at_ideation(world):
    out = build_gap_report(tenant_id=str(world["t"].id), project_id=str(world["project"].id),
                           object_store=world["store"])
    assert out["resume_phase"] == "ideation"
    assert out["blockers"] == ["ideation"]
    assert all(p["status"] == "missing" for p in out["phases"])
    # the report itself is a DRAFT internal artifact (published only on the gate, #5)
    report = Artifact.objects.get(id=out["gap_report_id"])
    assert report.type == "gap_report" and report.sor == "internal"
    assert report.status == Artifact.Status.DRAFT
    assert report.external_ref["resume_phase"] == "ideation"


def test_gap_report_picks_first_incomplete_phase_and_links_evidence(world):
    ideation = _published(world, "ideation")
    requirements = _published(world, "requirements")
    prd = _published(world, "prd")
    architecture = _draft(world, "architecture")           # registered, not yet adopted

    out = build_gap_report(tenant_id=str(world["t"].id), project_id=str(world["project"].id),
                           object_store=world["store"])
    status = {p["phase"]: p["status"] for p in out["phases"]}
    assert status["prd"] == "complete" and status["architecture"] == "partial"
    assert out["resume_phase"] == "architecture"           # first non-complete phase
    assert out["blockers"] == []                            # nothing Missing up to the resume

    # the gap report derives from every evidence artifact it weighed
    edges = {str(e.to_artifact_id) for e in TraceEdge.objects.filter(
        from_artifact_id=out["gap_report_id"])}
    assert edges == {ideation, requirements, prd, architecture}


def test_gap_report_intel_marks_implementation_assessed(world):
    tools.register_artifact(str(world["t"].id), str(world["project"].id), None,
                            phase="onboarding", type="codebase_intelligence",
                            sor="internal", content_hash="h")
    out = build_gap_report(tenant_id=str(world["t"].id), project_id=str(world["project"].id),
                           object_store=world["store"])
    assert out["intel_present"] is True
    status = {p["phase"]: p["status"] for p in out["phases"]}
    # code exists (intelligence proves it) → Implementation is assessed Complete, not Missing
    assert status["implementation"] == "complete"
    assert out["resume_phase"] == "ideation"               # docs still missing up front


# ── live Atlassian discovery (read-only; runs only with creds) ──────
def _live_config():
    from integrations.atlassian.config import AtlassianConfig

    cfg = AtlassianConfig.from_settings()
    return cfg if cfg.configured else None


@pytest.mark.skipif(_live_config() is None, reason="Atlassian creds not configured in .env")
@pytest.mark.django_db
def test_register_atlassian_live_imports_jira_backlog():
    from django.conf import settings

    from integrations.atlassian.config import AtlassianConfig
    from integrations.atlassian.tasktracker import AtlassianTaskTracker

    cfg = AtlassianConfig.from_settings()
    t = Tenant.objects.create(name="Brownfield Co")
    project = Project.objects.create(tenant=t, name="Existing")

    # B3-4: register_atlassian is Jira-only now (Confluence is linked + typed separately).
    out = register_atlassian(
        tenant_id=str(t.id), project_id=str(project.id),
        tasktracker=AtlassianTaskTracker(cfg, project=settings.JIRA_PROJECT_KEY),
        project_key=settings.JIRA_PROJECT_KEY,
    )
    assert out["tasktracker_connected"] is True
    # Any imported issue is a Jira-pointer WorkItem (register-don't-recreate; no body copied).
    for wi in WorkItem.objects.filter(project=project, external_ref__isnull=False):
        assert wi.external_ref["id"] and wi.external_ref["url"]
    # Idempotent: re-running imports nothing new.
    again = register_atlassian(
        tenant_id=str(t.id), project_id=str(project.id),
        tasktracker=AtlassianTaskTracker(cfg, project=settings.JIRA_PROJECT_KEY),
        project_key=settings.JIRA_PROJECT_KEY,
    )
    assert again["registered_issues"] == []


# ── Agent-driven onboarding (Slices 3–5) ─────────────────────────────
def test_write_gap_report_artifact_produces_readable_external_ref(world):
    """The ONE writer of the gap_report external_ref (shared by build_gap_report AND the agent
    path): read_gap_report must parse exactly what it writes."""
    tid, pid = str(world["t"].id), str(world["project"].id)
    phases = [{"phase": "ideation", "label": "1. Ideation", "status": "complete", "evidence": []},
              {"phase": "requirements", "label": "2. Requirements", "status": "missing",
               "evidence": []}]
    out = write_gap_report_artifact(
        tenant_id=tid, project_id=pid, phases=phases, resume="requirements",
        blockers=["requirements"], intel_present=False, object_store=world["store"])
    art = Artifact.objects.get(id=out["gap_report_id"])
    assert set(art.external_ref) == {
        "source", "resume_phase", "phases", "blockers", "intel_present"}
    assert set(art.external_ref["phases"][0]) >= {"phase", "label", "status", "evidence"}
    read = read_gap_report(tenant_id=tid, project_id=pid)
    assert read["resume_phase"] == "requirements"
    assert read["phases"][0]["assessment"] == "complete"


def test_submit_agent_gap_report_normalizes_and_reads_back(world):
    """The agent supplies a loose per-phase assessment; the tool coerces it to the full 12-phase
    canonical shape (omitted → missing, unknown phase ignored) and derives the resume phase."""
    tid, pid = str(world["t"].id), str(world["project"].id)
    submit_agent_gap_report(
        tenant_id=tid, project_id=pid,
        phases=[{"phase": "ideation", "status": "complete"},
                {"phase": "requirements", "status": "complete"},
                {"phase": "bogus_phase", "status": "complete"}],  # unknown → ignored
        resume_phase=None, blockers=None, object_store=world["store"])
    read = read_gap_report(tenant_id=tid, project_id=pid)
    assert len(read["phases"]) == 12                        # every lifecycle phase present
    by = {p["phase"]: p["assessment"] for p in read["phases"]}
    assert by["ideation"] == "complete" and by["requirements"] == "complete"
    assert by["prd"] == "missing"                           # omitted → missing
    assert read["resume_phase"] == "prd"                    # first non-complete
    assert "bogus_phase" not in by


def test_ensure_onboarding_work_item_idempotent_per_repo(world):
    from apps.projects.models import Repository

    tid, pid = str(world["t"].id), str(world["project"].id)
    a = ensure_onboarding_work_item(tenant_id=tid, project_id=pid)
    b = ensure_onboarding_work_item(tenant_id=tid, project_id=pid)
    assert a.id == b.id and a.type == WorkItem.Type.EPIC     # one repo-agnostic carrier
    repo = Repository.objects.create(tenant=world["t"], project=world["project"], name="acme/api")
    r1 = ensure_onboarding_work_item(tenant_id=tid, project_id=pid, repo_id=str(repo.id))
    r2 = ensure_onboarding_work_item(tenant_id=tid, project_id=pid, repo_id=str(repo.id))
    assert r1.id == r2.id and r1.id != a.id and r1.repo_id == repo.id  # a separate pinned carrier


def test_trigger_onboarding_agent_fails_closed_without_companion(world):
    """No companion enrolled → NoCompanionError (never a silent server-mode run that lacks the
    worktree + ask_human tools)."""
    tid, pid = str(world["t"].id), str(world["project"].id)
    with pytest.raises(NoCompanionError):
        trigger_onboarding_agent(
            tenant_id=tid, project_id=pid, agent_key="00_onboarding",
            triggered_by=str(world["po"].id), workflow_starter=lambda *a, **k: None)


def test_trigger_onboarding_agent_dispatches_companion_run_when_enrolled(world):
    from apps.registry.catalog import seed_agent_specs
    from apps.runs.enrollment import enroll_runner
    from apps.runs.models import Run

    tid, pid = str(world["t"].id), str(world["project"].id)
    seed_agent_specs(world["t"])
    enroll_runner(tenant_id=tid, project_id=pid, developer_id=str(world["po"].id), label="l")
    seen = {}
    res = trigger_onboarding_agent(
        tenant_id=tid, project_id=pid, agent_key="00_onboarding",
        triggered_by=str(world["po"].id),
        workflow_starter=lambda wf_id, inp: seen.update(inp=inp))
    assert seen["inp"].execution_mode == "companion"        # fail-closed to the companion
    carrier = WorkItem.objects.get(
        tenant=world["t"], project=world["project"], type=WorkItem.Type.EPIC)
    assert seen["inp"].work_item_id == str(carrier.id)      # ran against the carrier WI
    assert Run.objects.filter(id=res.run_id).exists()
