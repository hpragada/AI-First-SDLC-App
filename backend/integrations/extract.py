"""Static-file extraction (B4) — Word/Excel/PDF/Text/HTML → markdown text + images.

Best-effort **high-fidelity** extraction so an uploaded document becomes a reviewable internal
artifact body without losing content (onboarding_fidelity_plan Part A). Each adapter returns
``(markdown_body, images)`` where ``images`` is a list of ``{name, bytes, mime, page}`` for the
raster pictures embedded in the source — the ingest layer persists those and renders page images
so nothing pictorial is silently dropped. PDFs use PyMuPDF (text+layout, tables reconstructed as
markdown, embedded images); Word pulls headings/paragraphs/tables + embedded images; Excel gives
sheets-as-tables; plain text passes through; HTML is stripped to text with `<img>` captured. The
body is markdown (the artifact-body format used everywhere else).
"""
from __future__ import annotations

import base64
import io
import re
from html.parser import HTMLParser


class ExtractionError(Exception):
    """The file could not be extracted (unsupported type or corrupt content)."""


# One extracted raster image: name (stable-ish), raw bytes, mime, and source page (PDF) or None.
Image = dict


def _table_md(rows: list[list]) -> str:
    rows = [[("" if c is None else str(c)) for c in r] for r in rows]
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    out = [
        "| " + " | ".join(norm[0]) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for r in norm[1:]:
        out.append("| " + " | ".join(c.replace("|", "\\|").replace("\n", " ") for c in r) + " |")
    return "\n".join(out)


def _hardbreak(text: str) -> str:
    """Keep each visual line on its own line (markdown hard break = trailing two spaces) so the
    CommonMark preview doesn't collapse single newlines into one wall-of-text paragraph."""
    lines = [ln.rstrip() for ln in text.splitlines()]
    return "\n".join(f"{ln}  " if ln else "" for ln in lines).strip()


def _extract_docx(content: bytes) -> tuple[str, list[Image]]:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit())
            parts.append("#" * (int(level) if level else 2) + " " + text)
        else:
            parts.append(text)
    for table in doc.tables:
        md = _table_md([[c.text.strip() for c in row.cells] for row in table.rows])
        if md:
            parts.append(md)
    # Embedded images live as document relationships (Part A: don't drop pictures).
    images: list[Image] = []
    for rel in doc.part.rels.values():
        try:
            if "image" not in rel.reltype:
                continue
            part = rel.target_part
            ct = getattr(part, "content_type", "") or "image/png"
            ext = (ct.split("/")[-1].split("+")[0] or "png").lower()
            images.append({"name": f"docx-img{len(images) + 1}.{ext}", "bytes": part.blob,
                           "mime": ct, "page": None})
        except Exception:  # noqa: BLE001 — a bad rel must not fail the whole extraction
            continue
    return "\n\n".join(parts), images


def _extract_xlsx(content: bytes) -> tuple[str, list[Image]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        md = _table_md(rows)
        if md:
            parts.append(f"## {ws.title}\n\n{md}")
    return "\n\n".join(parts), []


def _table_is_clean(rows: list[list]) -> bool:
    """A confidence gate on a detected PDF table. Table *detection* on complex bordered PDFs is
    unreliable — it invents phantom empty columns and duplicates cells. Emitting a garbled grid
    as if authoritative is worse than clean text for a system meant to be trusted, so we only
    accept a table that is well-formed: >=2 rows, >=2 columns, and no fully-empty column (a
    fully-empty column is the tell-tale of mis-detection). Rejected tables fall back to the page's
    plain text; the page image (A2) + the vision step (Part B) carry the true table faithfully."""
    if len(rows) < 2:
        return False
    width = max((len(r) for r in rows), default=0)
    if width < 2:
        return False
    for c in range(width):
        col_has_text = any(str((r[c] if c < len(r) else "") or "").strip() for r in rows)
        if not col_has_text:
            return False
    return True


def _pdf_page_markdown(page, pymupdf) -> str:
    """One PDF page → markdown: non-table text (in reading order) interleaved with any *cleanly
    detected* tables reconstructed as markdown. A clean table's text region is removed from the
    surrounding text so it isn't duplicated; a low-confidence table is left as plain text."""
    try:
        detected = list(page.find_tables().tables)
    except Exception:  # noqa: BLE001
        detected = []

    clean: list[tuple] = []  # (table, rows) that passed the confidence gate
    for t in detected:
        try:
            rows = [list(r) for r in t.extract()]
        except Exception:  # noqa: BLE001
            continue
        if _table_is_clean(rows):
            clean.append((t, rows))

    table_rects = []
    for t, _rows in clean:
        try:
            table_rects.append(pymupdf.Rect(t.bbox))
        except Exception:  # noqa: BLE001
            pass

    def _in_table(bbox) -> bool:
        r = pymupdf.Rect(bbox)
        if r.is_empty:
            return False
        for tr in table_rects:
            inter = r & tr
            if not inter.is_empty and abs(inter) >= 0.5 * abs(r):
                return True
        return False

    items: list[tuple[float, str]] = []  # (y-position, markdown) — sorted into reading order
    for b in page.get_text("blocks"):
        x0, y0, x1, y1, txt, _bno, btype = b
        if btype != 0 or not str(txt).strip():
            continue
        if _in_table((x0, y0, x1, y1)):
            continue
        md = _hardbreak(str(txt))
        if md:
            items.append((y0, md))
    for t, rows in clean:
        md = _table_md(rows)
        if md:
            items.append((pymupdf.Rect(t.bbox).y0, md))
    items.sort(key=lambda it: it[0])
    return "\n\n".join(md for _y, md in items).strip()


def _extract_pdf(content: bytes) -> tuple[str, list[Image]]:
    import pymupdf

    parts: list[str] = []
    images: list[Image] = []
    with pymupdf.open(stream=content, filetype="pdf") as doc:
        for pno, page in enumerate(doc, start=1):
            block = _pdf_page_markdown(page, pymupdf)
            if block:
                parts.append(block)
            for img in page.get_images(full=True):
                xref = img[0]
                try:
                    base = doc.extract_image(xref)
                except Exception:  # noqa: BLE001
                    continue
                ext = (base.get("ext") or "png").lower()
                images.append({"name": f"page{pno}-img{xref}.{ext}", "bytes": base["image"],
                               "mime": f"image/{ext}", "page": pno})
    return "\n\n".join(parts), images


def pdf_page_pngs(content: bytes, *, dpi: int = 130) -> list[Image]:
    """Render each PDF page to a PNG — the faithful visual of the upload, for the 'see your
    original' viewer (A3) and the vision step (Part B). Best-effort per page; `dpi` kept modest
    so storage + downstream vision cost stay sane."""
    import pymupdf

    out: list[Image] = []
    with pymupdf.open(stream=content, filetype="pdf") as doc:
        for pno, page in enumerate(doc, start=1):
            try:
                pix = page.get_pixmap(dpi=dpi)
                out.append({"name": f"page{pno}.png", "bytes": pix.tobytes("png"),
                            "mime": "image/png", "page": pno})
            except Exception:  # noqa: BLE001 — one bad page must not lose the rest
                continue
    return out


def _extract_txt(content: bytes) -> tuple[str, list[Image]]:
    """Plain text is already the artifact-body format — decode and return as-is."""
    return content.decode("utf-8", "replace").strip(), []


class _HTMLTextExtractor(HTMLParser):
    """Strip HTML to readable text: skip script/style, keep headings (as markdown #), list items
    (as '- '), block boundaries (as blank lines), and capture `<img>` (Part A: keep pictures —
    data-URI images carry bytes; external ones keep their src). Stdlib only."""

    _SKIP = {"script", "style", "head", "noscript", "template"}
    _HEADINGS = {"h1", "h2", "h3", "h4", "h5", "h6"}
    _BLOCK = {"p", "div", "section", "article", "header", "footer", "br",
              "ul", "ol", "table", "tr", "blockquote", "pre"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.images: list[Image] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag in self._HEADINGS:
            self.parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag == "li":
            self.parts.append("\n- ")
        elif tag == "img":
            self._capture_img(dict(attrs))
        elif tag in self._BLOCK:
            self.parts.append("\n\n")

    def _capture_img(self, attrs: dict) -> None:
        src = attrs.get("src") or ""
        alt = attrs.get("alt") or ""
        n = len(self.images) + 1
        m = re.match(r"data:(image/[\w.+-]+);base64,(.*)$", src, re.DOTALL)
        if m:
            mime = m.group(1)
            try:
                raw = base64.b64decode(m.group(2), validate=False)
            except Exception:  # noqa: BLE001
                raw = b""
            if raw:
                ext = mime.split("/")[-1].split("+")[0] or "png"
                self.images.append({"name": f"html-img{n}.{ext}", "bytes": raw,
                                    "mime": mime, "page": None})
        elif src:
            # External image — keep the reference (no bytes to persist without a fetch).
            self.images.append({"name": alt or f"html-img{n}", "src": src,
                                "mime": None, "bytes": None, "page": None})

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._HEADINGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = " ".join(data.split())
        if text:
            self.parts.append(text + " ")


def _extract_html(content: bytes) -> tuple[str, list[Image]]:
    parser = _HTMLTextExtractor()
    parser.feed(content.decode("utf-8", "replace"))
    parser.close()
    lines = [ln.rstrip() for ln in "".join(parser.parts).splitlines()]
    out = "\n".join(lines)
    # Inline tags (<b>, <a>, …) split text into chunks each padded with a trailing space,
    # which leaves a gap before punctuation ("login ."). Close those gaps.
    out = re.sub(r"[ \t]+([.,;:!?)])", r"\1", out)
    while "\n\n\n" in out:
        out = out.replace("\n\n\n", "\n\n")
    return out.strip(), parser.images


def _kind(mime: str, ext: str) -> str | None:
    mime = (mime or "").lower()
    if ext == "docx" or "wordprocessingml" in mime:
        return "docx"
    if ext == "xlsx" or "spreadsheetml" in mime:
        return "xlsx"
    if ext == "pdf" or mime == "application/pdf":
        return "pdf"
    if ext in ("html", "htm") or mime == "text/html":
        return "html"
    if ext == "txt" or mime == "text/plain":
        return "txt"
    return None


def extract(*, filename: str, mime: str, content: bytes) -> tuple[str, str, list[Image]]:
    """Return ``(title, markdown_body, images)`` for a supported static file. ``images`` holds the
    embedded raster pictures ({name, bytes, mime, page}); it's empty for text-only sources."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    kind = _kind(mime, ext)
    if kind is None:
        raise ExtractionError(f"unsupported file type: {filename!r} ({mime!r})")
    extractor = {
        "docx": _extract_docx, "xlsx": _extract_xlsx, "pdf": _extract_pdf,
        "txt": _extract_txt, "html": _extract_html,
    }[kind]
    try:
        body, images = extractor(content)
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001 — surface any parse failure as ExtractionError
        raise ExtractionError(f"failed to extract {filename!r}: {exc}") from exc
    title = filename.rsplit(".", 1)[0] if "." in filename else filename
    return title, body, images
